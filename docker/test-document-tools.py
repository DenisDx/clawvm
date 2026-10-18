#!/usr/bin/env python3
"""Exercise document tools on generated fixtures without network or user files."""

import subprocess
import tempfile
import unittest
from pathlib import Path

import fitz
import pdfplumber
from bs4 import BeautifulSoup
from docx import Document
from openpyxl import Workbook, load_workbook
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader


def run(*args: str) -> str:
    """Run a command with string arguments; return stdout or raise on failure."""
    return subprocess.run(args, check=True, text=True, stdout=subprocess.PIPE).stdout


class DocumentToolsTest(unittest.TestCase):
    """Verify PDF, OCR, graphics, office, HTML, and mail command availability."""

    @classmethod
    def setUpClass(cls) -> None:
        """Create temporary PDF fixtures in writable home; register cleanup."""
        temporary = tempfile.TemporaryDirectory(prefix=".document-smoke-", dir=Path.home())
        cls.addClassCleanup(temporary.cleanup)
        cls.directory = Path(temporary.name)
        cls.pdf = cls.directory / "source.pdf"
        cls.text = "Document tools smoke test 123"
        with fitz.open() as document:
            page = document.new_page()
            page.insert_text((72, 72), cls.text)
            document.save(cls.pdf)

    def test_pdf_text(self) -> None:
        """Extract the generated text through all four PDF readers."""
        self.assertIn(self.text, run("pdftotext", str(self.pdf), "-"))
        self.assertIn(self.text, PdfReader(self.pdf).pages[0].extract_text())
        with pdfplumber.open(self.pdf) as document:
            self.assertIn(self.text, document.pages[0].extract_text())
        with fitz.open(self.pdf) as document:
            self.assertIn(self.text, document[0].get_text())

    def test_pdf_render_and_rewrite(self) -> None:
        """Rasterize, linearize, and rewrite a PDF; verify the resulting files."""
        raster = self.directory / "raster"
        run("pdftoppm", "-singlefile", "-r", "72", "-png", str(self.pdf), str(raster))
        with Image.open(raster.with_suffix(".png")) as image:
            self.assertGreater(image.width, 500)
            self.assertGreater(image.height, 700)
        linear = self.directory / "linear.pdf"
        run("qpdf", "--linearize", str(self.pdf), str(linear))
        run("qpdf", "--check", str(linear))
        run("qpdf", "--check-linearization", str(linear))
        rewritten = self.directory / "rewritten.pdf"
        run("gs", "-q", "-dSAFER", "-dBATCH", "-dNOPAUSE", "-sDEVICE=pdfwrite",
            f"-sOutputFile={rewritten}", str(self.pdf))
        self.assertIn(self.text, PdfReader(rewritten).pages[0].extract_text())

    def test_ocr_languages(self) -> None:
        """Recognize rendered English, Russian, and Thai fixtures individually."""
        languages = run("tesseract", "--list-langs").splitlines()
        samples = (
            ("eng", "DOCUMENT TOOLS TEST", "NotoSans-Regular.ttf"),
            ("rus", "\u041f\u0440\u043e\u0432\u0435\u0440\u043a\u0430 "
             "\u0434\u043e\u043a\u0443\u043c\u0435\u043d\u0442\u043e\u0432",
             "NotoSans-Regular.ttf"),
            ("tha", "\u0e20\u0e32\u0e29\u0e32\u0e44\u0e17\u0e22", "NotoSansThai-Regular.ttf"),
        )
        for language, text, font_name in samples:
            with self.subTest(language=language):
                self.assertIn(language, languages)
                font = ImageFont.truetype(f"/usr/share/fonts/truetype/noto/{font_name}", 64)
                image = Image.new("RGB", (1200, 180), "white")
                ImageDraw.Draw(image).text((30, 25), text, font=font, fill="black")
                path = self.directory / f"ocr-{language}.png"
                image.save(path)
                recognized = run("tesseract", str(path), "stdout", "-l", language, "--psm", "6")
                self.assertIn("".join(text.split()), "".join(recognized.split()))

    def test_graphics_and_documents(self) -> None:
        """Round-trip DOCX/XLSX, parse HTML, and convert Markdown and SVG."""
        docx = self.directory / "source.docx"
        document = Document()
        document.add_paragraph(self.text)
        document.save(docx)
        self.assertEqual(Document(docx).paragraphs[0].text, self.text)
        xlsx = self.directory / "source.xlsx"
        workbook = Workbook()
        workbook.active["A1"] = self.text
        workbook.save(xlsx)
        loaded = load_workbook(xlsx)
        self.assertEqual(loaded.active["A1"].value, self.text)
        loaded.close()
        self.assertEqual(BeautifulSoup(f"<p>{self.text}</p>", "lxml").p.text, self.text)
        markdown = self.directory / "source.md"
        markdown.write_text(self.text, encoding="utf-8")
        converted = self.directory / "converted.docx"
        run("pandoc", str(markdown), "-o", str(converted))
        self.assertEqual(Document(converted).paragraphs[0].text, self.text)
        svg = self.directory / "source.svg"
        svg.write_text('<svg xmlns="http://www.w3.org/2000/svg" width="120" height="60">'
                       '<rect width="120" height="60" fill="red"/></svg>', encoding="utf-8")
        png = self.directory / "svg.png"
        run("rsvg-convert", str(svg), "-o", str(png))
        with Image.open(png) as image:
            self.assertEqual(image.size, (120, 60))

    def test_mail_and_hex_tools(self) -> None:
        """Check mail clients without contacting servers, plus a real hexdump."""
        self.assertIn("msmtp", run("msmtp", "--version"))
        self.assertIn("isync", run("mbsync", "--version"))
        self.assertIn("25504446", run("xxd", "-l", "4", "-p", str(self.pdf)))


if __name__ == "__main__":
    unittest.main(verbosity=2)
